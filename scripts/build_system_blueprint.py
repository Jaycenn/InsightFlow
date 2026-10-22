"""Build the editable InsightFlow system blueprint from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "InsightFlow-System-Blueprint.md"
OUTPUT = ROOT / "docs" / "InsightFlow-System-Blueprint.docx"
ASSETS = ROOT / "docs" / "assets" / "system-blueprint"

INK = "152530"
BLUE = "173A52"
TEAL = "20776F"
LIGHT = "EAF3F3"
PALE = "F4F8FA"
LINE = "C9D7DE"
WHITE = "FFFFFF"


def font(size: int = 35, bold: bool = False) -> ImageFont.FreeTypeFont:
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / filename), size)


def wrap(draw: ImageDraw.ImageDraw, text: str, width: int, face: ImageFont.FreeTypeFont) -> str:
    lines: list[str] = []
    line = ""
    for word in text.split():
        trial = (line + " " + word).strip()
        if draw.textbbox((0, 0), trial, font=face)[2] <= width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return "\n".join(lines)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((*start, *end), fill=f"#{TEAL}", width=6)
    x, y = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x, y), (x - 18 * direction, y - 12), (x - 18 * direction, y + 12)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x, y), (x - 12, y - 18 * direction), (x + 12, y - 18 * direction)]
    draw.polygon(points, fill=f"#{TEAL}")


def card(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], label: str, *, dark: bool = False) -> None:
    fill = BLUE if dark else PALE
    outline = BLUE if dark else LINE
    draw.rounded_rectangle(box, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=3)
    face = font(34 if len(label) < 34 else 31, True)
    text = wrap(draw, label, box[2] - box[0] - 38, face)
    bounds = draw.multiline_textbbox((0, 0), text, font=face, spacing=7, align="center")
    tx = (box[0] + box[2] - (bounds[2] - bounds[0])) // 2
    ty = (box[1] + box[3] - (bounds[3] - bounds[1])) // 2 - 5
    draw.multiline_text((tx, ty), text, font=face, fill=f"#{WHITE if dark else INK}", spacing=7, align="center")


def flow_diagram(filename: str, rows: list[list[str]], dark_first: bool = True) -> Path:
    width = 1800
    max_cols = max(len(row) for row in rows)
    h = 120 + 180 * len(rows) + 90 * (len(rows) - 1)
    image = Image.new("RGB", (width, h), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    box_h = 154
    positions: list[list[tuple[int, int, int, int]]] = []
    for row_i, row in enumerate(rows):
        margin = 68
        gap = 42
        box_w = (width - 2 * margin - gap * (len(row) - 1)) // len(row)
        y = 55 + row_i * 265
        row_boxes = []
        for col_i, label in enumerate(row):
            x = margin + col_i * (box_w + gap)
            box = (x, y, x + box_w, y + box_h)
            card(draw, box, label, dark=dark_first and row_i == 0 and col_i == 0)
            row_boxes.append(box)
        positions.append(row_boxes)
        for left, right in zip(row_boxes, row_boxes[1:]):
            arrow(draw, (left[2] + 8, y + box_h // 2), (right[0] - 8, y + box_h // 2))
    # Each row reads from left to right; rows continue from top to bottom.
    path = ASSETS / filename
    image.save(path)
    return path


def er_diagram() -> Path:
    image = Image.new("RGB", (1800, 720), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    boxes = {
        "User": (70, 250, 340, 395),
        "Business": (510, 250, 810, 395),
        "Dataset": (1020, 55, 1340, 200),
        "Product": (1020, 250, 1340, 395),
        "Customer": (1020, 445, 1340, 590),
        "Sale line": (1490, 250, 1760, 395),
    }
    arrow(draw, (350, 323), (500, 323))
    for target in ("Dataset", "Product", "Customer"):
        b = boxes[target]
        arrow(draw, (820, 323), (b[0] - 10, (b[1] + b[3]) // 2))
    arrow(draw, (1350, 128), (1480, 310))
    arrow(draw, (1350, 323), (1480, 323))
    arrow(draw, (1350, 517), (1480, 337))
    for name, box in boxes.items():
        card(draw, box, name, dark=name == "Business")
    path = ASSETS / "entity-relationships.png"
    image.save(path)
    return path


def architecture_diagram() -> Path:
    image = Image.new("RGB", (1800, 660), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    browser = (55, 70, 330, 220)
    web = (420, 70, 775, 220)
    api = (865, 70, 1185, 220)
    database = (1470, 70, 1760, 220)
    arrow(draw, (340, 145), (410, 145))
    arrow(draw, (785, 145), (855, 145))
    arrow(draw, (1195, 145), (1460, 145))
    card(draw, browser, "Browser", dark=True)
    card(draw, web, "React and TanStack Query")
    card(draw, api, "FastAPI REST API")
    card(draw, database, "PostgreSQL")
    modules = [
        ((650, 390, 960, 540), "Authentication"),
        ((1020, 390, 1330, 540), "Import service"),
        ((1390, 390, 1730, 540), "Analytics service"),
    ]
    for box, label in modules:
        arrow(draw, (1025, 230), ((box[0] + box[2]) // 2, box[1] - 10))
        card(draw, box, label)
    path = ASSETS / "implemented-architecture.png"
    image.save(path)
    return path


def dataset_diagram() -> Path:
    image = Image.new("RGB", (1800, 590), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    uploaded = (80, 90, 460, 250)
    imported = (710, 90, 1090, 250)
    removed = (1370, 90, 1740, 250)
    validation = (80, 370, 460, 530)
    draw.line((270, 78, 270, 35, 1555, 35, 1555, 78), fill=f"#{TEAL}", width=6, joint="curve")
    draw.polygon([(1555, 90), (1543, 70), (1567, 70)], fill=f"#{TEAL}")
    arrow(draw, (470, 170), (700, 170))
    arrow(draw, (1100, 170), (1360, 170))
    arrow(draw, (270, 260), (270, 360))
    arrow(draw, (470, 450), (700, 220))
    card(draw, uploaded, "Uploaded dataset", dark=True)
    card(draw, imported, "Imported dataset")
    card(draw, removed, "Deleted dataset")
    card(draw, validation, "Preview and validation")
    path = ASSETS / "dataset-lifecycle.png"
    image.save(path)
    return path


def production_diagram() -> Path:
    image = Image.new("RGB", (1800, 760), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    browser = (50, 70, 330, 220)
    web = (405, 70, 705, 220)
    api = (790, 70, 1090, 220)
    database = (1450, 70, 1750, 220)
    arrow(draw, (340, 145), (395, 145))
    arrow(draw, (715, 145), (780, 145))
    arrow(draw, (1100, 145), (1440, 145))
    billing = (50, 460, 360, 610)
    webhook = (450, 460, 760, 610)
    storage = (850, 460, 1160, 610)
    recovery = (1440, 460, 1750, 610)
    arrow(draw, (370, 535), (440, 535))
    arrow(draw, (770, 535), (940, 230))
    arrow(draw, (940, 230), (1005, 450))
    arrow(draw, (1600, 230), (1600, 450))
    for box, label, dark in [
        (browser, "Customer browser", True), (web, "HTTPS web app", False),
        (api, "HTTPS API", False), (database, "Managed PostgreSQL", False),
        (billing, "Billing provider", False), (webhook, "Verified webhook", False),
        (storage, "Storage and jobs if needed", False), (recovery, "Backups and restore", False),
    ]:
        card(draw, box, label, dark=dark)
    path = ASSETS / "production-topology.png"
    image.save(path)
    return path


def draw_diagrams() -> list[Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    return [
        flow_diagram("customer-journey.png", [
            ["Create account", "Create workspace", "Upload sales file", "Preview and map"],
            ["Validate rows", "Import sales", "Explore dashboards", "Upload next export"],
        ]),
        architecture_diagram(),
        flow_diagram("data-flow.png", [
            ["Customer file", "Limit and parse", "Preview and map", "Validate"],
            ["Database transaction", "Business sales", "Filtered SQL", "Dashboards"],
        ]),
        flow_diagram("local-topology.png", [["Browser", "Nginx web frontend", "FastAPI backend", "PostgreSQL 16"]]),
        er_diagram(),
        dataset_diagram(),
        production_diagram(),
        flow_diagram("billing-lifecycle.png", [
            ["Choose plan", "Hosted checkout", "Signed webhook", "Save subscription"],
            ["Check entitlement", "Use paid service", "Renew or cancel", "Update access"],
        ]),
    ]


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top: int = 100, side: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("bottom", top), ("left", side), ("right", side)):
        el = margins.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def add_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif part.startswith("[") and "](" in part:
            match = re.match(r"\[([^]]+)\]\(([^)]+)\)", part)
            if match:
                paragraph.add_run(f"{match.group(1)} ({match.group(2)})")
        else:
            paragraph.add_run(part)


def add_table(doc: Document, source_lines: list[str]) -> None:
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in source_lines]
    rows.pop(1)  # Markdown separator
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = 6.92
    if len(rows[0]) == 4 and rows[0][0] == "Canonical field":
        proportions = [0.19, 0.12, 0.27, 0.42]
    elif len(rows[0]) == 4 and rows[0][0] == "Method":
        proportions = [0.11, 0.39, 0.38, 0.12]
    elif len(rows[0]) == 4:
        proportions = [0.17, 0.31, 0.38, 0.14]
    elif len(rows[0]) == 3:
        proportions = [0.22, 0.32, 0.46]
    elif len(rows[0]) == 2:
        proportions = [0.29, 0.71]
    else:
        proportions = [1 / len(rows[0])] * len(rows[0])
    for i, data in enumerate(rows):
        for j, value in enumerate(data):
            cell = table.cell(i, j)
            cell.width = Inches(total * proportions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            if i == 0:
                shade(cell, BLUE)
            elif i % 2 == 0:
                shade(cell, PALE)
            para = cell.paragraphs[0]
            para.style = doc.styles["Normal"]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.12
            add_runs(para, value)
            for run in para.runs:
                run.font.size = Pt(8.3 if len(rows[0]) < 4 else 7.8)
                if i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        if i == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.17
    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(11)
    for name, size, before, after in (("Heading 1", 15, 17, 8), ("Heading 2", 11.5, 12, 5), ("Heading 3", 10, 10, 4)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(9.4)
        style.paragraph_format.space_after = Pt(3)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("InsightFlow System Blueprint  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 105, 115)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build() -> None:
    diagrams = draw_diagrams()
    doc = Document()
    configure(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    diagram_index = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("```mermaid"):
            while i < len(lines) and not lines[i].startswith("```") or (i < len(lines) and lines[i] == "```mermaid"):
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p.add_run().add_picture(str(diagrams[diagram_index]), width=Inches(6.75))
            diagram_index += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.3)
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s|:-]+\|$", lines[i + 1]):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                doc.add_paragraph(heading.group(2), style="Title")
            else:
                doc.add_paragraph(heading.group(2), style=f"Heading {min(level - 1, 3)}")
            i += 1
            continue
        ordered = re.match(r"^\d+\.\s+(.+)$", line)
        if ordered:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, ordered.group(1))
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, bullet.group(1))
            i += 1
            continue
        p = doc.add_paragraph(style="Normal")
        add_runs(p, line.replace("  ", " "))
        i += 1
    if diagram_index != len(diagrams):
        raise RuntimeError(f"Used {diagram_index} diagrams, expected {len(diagrams)}")
    doc.core_properties.title = "InsightFlow System Blueprint"
    doc.core_properties.subject = "Current system design and hosted subscription plan"
    doc.core_properties.author = "InsightFlow"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
